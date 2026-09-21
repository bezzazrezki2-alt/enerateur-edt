import sys
import subprocess
import tkinter as tk
from tkinter import messagebox

try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document

from docx.shared import Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def generer_document():
    # --- 1. RÉCUPÉRATION DES DONNÉES ---
    nom = entry_nom.get().strip() or "رزقي محمد رزقي"
    grade = entry_grade.get().strip() or "مستشار الشباب"
    activite = entry_activite.get().strip() or "سمعي بصري"
    fonction = entry_fonction.get().strip() or "سمعي بصري"
    activite_sup = entry_activite_sup.get().strip()
    
    try:
        total_heures = int(entry_heures.get())
    except ValueError:
        messagebox.showerror("Erreur", "Veuillez entrer un nombre d'heures valide.")
        return

    jours_choisis = [jour for jour, var in vars_jours.items() if var.get()]
    if not jours_choisis:
        messagebox.showerror("Erreur", "Veuillez cocher au moins un jour de travail.")
        return

    # --- 2. LOGIQUE DE RÉPARTITION ---
    jours_semaine = ["الأحد", "الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت"]
    schedule_final = {jour: [] for jour in jours_semaine}

    heures_par_jour = total_heures // len(jours_choisis)
    heures_restantes = total_heures % len(jours_choisis)

    for i, jour in enumerate(jours_choisis):
        h_a_placer = heures_par_jour + (1 if i < heures_restantes else 0)
        
        type_horaire = vars_horaires[jour].get()
        
        if type_horaire == "Matinée":
            # Uniquement le matin (en évitant l'heure entre midi et 13h)
            pool = ["10-09", "11-10", "12-11"]
        elif type_horaire == "Après-midi":
            # Uniquement l'après-midi (dès 13h ou dès 15h pour le vendredi)
            if jour == "الجمعة":
                pool = ["16-15", "17-16", "18-17", "19-18", "20-19", "21-20"]
            else:
                pool = ["14-13", "15-14", "16-15", "17-16", "18-17", "19-18", "20-19", "21-20"]
        else: # Mixte (Matin et Après-midi)
            if jour == "الجمعة":
                pool = ["10-09", "11-10", "12-11", "16-15", "17-16", "18-17", "19-18", "20-19"]
            else:
                pool = ["10-09", "11-10", "12-11", "14-13", "15-14", "16-15", "17-16", "18-17", "19-18", "20-19"]
            
        schedule_final[jour] = pool[:h_a_placer]

    # --- 3. CRÉATION DU DOCUMENT WORD ---
    doc = Document()
    
    section = doc.sections[-1]
    new_w, new_h = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_w
    section.page_height = new_h
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("الجمهورية الجزائرية الديمقراطية الشعبية\nوزارة الشباب\n")
    run.bold = True
    
    table_header = doc.add_table(rows=1, cols=3)
    table_header.autofit = True
    c1, c2, c3 = table_header.rows[0].cells
    
    c3.text = "ولاية الجزائر\nمديرية الشباب والرياضة و الترفيه\nديوان مؤسسات الشباب\nالمؤسسة : دار الشباب محمد أوبنير حميدوقن بني مسوس"
    c3.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    c2.text = f"الاسم و اللقب: {nom}\nالرتبة: {grade}\nالحجم الساعي القانوني: {total_heures} سا"
    c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    texte_gauche = f"النشاط الأساسي: {activite}\n\nالوظيفة: {fonction}"
    if activite_sup:
        texte_gauche += f"\n\nالنشاط الإضافي: {activite_sup}"
        
    c1.text = texte_gauche
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    p_title = doc.add_paragraph("\n")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("جدول التوقيت الأسبوعي للموسم 2026/2027")
    run_title.bold = True
    run_title.font.size = Pt(16)
    
    colonnes = [
        "21-20", "20-19", "19-18", "18-17", "17-16", "16-15", "15-14",
        "14-13", "13-12", "12-11", "11-10", "10-09", "09-08", "الأيام"
    ]
    
    table = doc.add_table(rows=8, cols=14)
    table.style = 'Table Grid'
    
    for col_idx, col_name in enumerate(colonnes):
        cell = table.cell(0, col_idx)
        cell.text = col_name
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    activite_cell_text = "\n".join(activite.split()) 

    for row_idx, jour in enumerate(jours_semaine, start=1):
        cell_day = table.cell(row_idx, 13)
        cell_day.text = jour
        cell_day.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        for col_idx, col_name in enumerate(colonnes[:-1]):
            if col_name in schedule_final[jour]:
                cell = table.cell(row_idx, col_idx)
                cell.text = activite_cell_text
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'E6D8E7')
                tcPr.append(shd)

    p_footer = doc.add_paragraph("\n\n")
    p_footer.add_run("مفتش المقاطعة أو مستشار الشباب                                            مدير المؤسسة                                            المعني")
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    nom_fichier = f"Emploi_du_temps_{nom.replace(' ', '_')}_{total_heures}h.docx"
    doc.save(nom_fichier)
    
    messagebox.showinfo("Succès", f"L'emploi du temps a été généré avec succès !\nFichier : {nom_fichier}")

# --- 4. INTERFACE GRAPHIQUE (TKINTER) ---
fenetre = tk.Tk()
fenetre.title("Générateur d'Emploi du Temps")
fenetre.geometry("550x700")
fenetre.configure(padx=20, pady=20)

tk.Label(fenetre, text="Informations de l'Employé", font=("Arial", 12, "bold")).grid(row=0, column=0, columnspan=2, pady=(0, 10), sticky="w")

tk.Label(fenetre, text="Nom et Prénom :").grid(row=1, column=0, sticky="w", pady=2)
entry_nom = tk.Entry(fenetre, width=30)
entry_nom.insert(0, "رزقي محمد رزقي")
entry_nom.grid(row=1, column=1, pady=2, sticky="w")

tk.Label(fenetre, text="Grade (الرتبة) :").grid(row=2, column=0, sticky="w", pady=2)
entry_grade = tk.Entry(fenetre, width=30)
entry_grade.insert(0, "مستشار الشباب")
entry_grade.grid(row=2, column=1, pady=2, sticky="w")

tk.Label(fenetre, text="Fonction (الوظيفة) :").grid(row=3, column=0, sticky="w", pady=2)
entry_fonction = tk.Entry(fenetre, width=30)
entry_fonction.insert(0, "سمعي بصري")
entry_fonction.grid(row=3, column=1, pady=2, sticky="w")

tk.Label(fenetre, text="Activité Principale :").grid(row=4, column=0, sticky="w", pady=2)
entry_activite = tk.Entry(fenetre, width=30)
entry_activite.insert(0, "سمعي بصري")
entry_activite.grid(row=4, column=1, pady=2, sticky="w")

tk.Label(fenetre, text="Activité Suppl. (Optionnel) :").grid(row=5, column=0, sticky="w", pady=2)
entry_activite_sup = tk.Entry(fenetre, width=30)
entry_activite_sup.grid(row=5, column=1, pady=2, sticky="w")

tk.Label(fenetre, text="Configuration des Heures", font=("Arial", 12, "bold")).grid(row=6, column=0, columnspan=2, pady=(20, 10), sticky="w")

tk.Label(fenetre, text="Volume horaire :").grid(row=7, column=0, sticky="w", pady=2)
entry_heures = tk.Spinbox(fenetre, from_=1, to=100, width=10)
entry_heures.delete(0, "end")
entry_heures.insert(0, "20")
entry_heures.grid(row=7, column=1, pady=2, sticky="w")

tk.Label(fenetre, text="Jours travaillés :").grid(row=8, column=0, sticky="nw", pady=5)
frame_jours = tk.Frame(fenetre)
frame_jours.grid(row=8, column=1, sticky="w")

jours = ["السبت", "الأحد", "الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة"]
jours_par_defaut = ["السبت", "الأحد", "الثلاثاء", "الخميس", "الجمعة"]
vars_jours = {}
vars_horaires = {}

for i, jour in enumerate(jours):
    var_check = tk.BooleanVar(value=(jour in jours_par_defaut))
    vars_jours[jour] = var_check
    tk.Checkbutton(frame_jours, text=jour, variable=var_check).grid(row=i, column=0, sticky="w")
    
    if jour == "الجمعة":
        horaire_defaut = "Après-midi"
    else:
        horaire_defaut = "Après-midi" if jour in ["الأحد", "الثلاثاء"] else "Matinée"
        
    var_horaire = tk.StringVar(value=horaire_defaut)
    vars_horaires[jour] = var_horaire
    
    opt = tk.OptionMenu(frame_jours, var_horaire, "Matinée", "Après-midi", "Mixte")
    opt.config(width=15, font=("Arial", 9))
    opt.grid(row=i, column=1, padx=10, pady=2, sticky="w")

tk.Button(fenetre, text="Générer l'Emploi du Temps", bg="#4CAF50", fg="white", font=("Arial", 11, "bold"), command=generer_document).grid(row=9, column=0, columnspan=2, pady=30, sticky="we")

fenetre.mainloop()